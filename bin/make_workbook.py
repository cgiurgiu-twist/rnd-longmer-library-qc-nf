#!/usr/bin/env python3
"""
make_workbook.py — the multi-sheet Excel workbook and the plain-text run summary.

One sheet per question a reviewer will ask: headline metrics with their denominators,
the calibration controls, read length by class, the block-deletion detail, the
full-length definition ladder, screening depth, and the per-variant / per-well tables.

The README sheet carries the denominators and bounds, so the workbook is readable
without the report.
"""
import argparse, csv, json, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

AP = argparse.ArgumentParser()
AP.add_argument("--metrics-dir", default=".")
AP.add_argument("--meta", required=True)
AP.add_argument("--outdir", default=".")
AP.add_argument("--name", default="library")
AP.add_argument("--microhomology", default=None)
A = AP.parse_args()
D = A.metrics_dir.rstrip("/") + "/"; O = A.outdir.rstrip("/") + "/"
os.makedirs(O, exist_ok=True)
J = lambda n: json.load(open(D + n))
core, ext, der = J("metrics_core.json"), J("metrics_extra.json"), J("derived.json")
BD, SC, CT = J("blockdel_summary.json"), J("blockdel_scope.json"), J("controls.json")
ORD = sorted(core["per_sublib"])
meta = json.load(open(A.meta))
ov, ps, U, REC, bd = core["overall"], core["per_sublib"], der["uniformity"], der["recoverability"], der["blockdel"]
dsg = der.get("design", {})
DEL = {o: SC["per_sublib_any_cls"].get(o, {}).get("pct") for o in ORD}
tm = {o: ps[o]["within_well_pct_of_assigned"] + ps[o]["cross_well_pct_of_assigned"] + ps[o]["cross_sublib_pct_of_assigned"] for o in ORD}
TMall = ov["within_well_pct_of_assigned"] + ov["cross_well_pct_of_assigned"] + ov["cross_sublib_pct_of_assigned"]
sm = ext.get("single_molecule_pairing", {})
def fmt3(d, *keys):
    """Three numbers or the reason they are absent — never a crash on a sub-library the
    subsample did not reach."""
    if not all(d.get(k) is not None for k in keys):
        return d.get("note", "not computed")
    return " / ".join(f"{d[k]:.2f}" for k in keys)


wb = Workbook(); HD = Font(bold=True, color="FFFFFF"); FILL = PatternFill("solid", fgColor="365F91")
def sheet(name, hdr, rows, widths=None):
    ws = wb.create_sheet(name); ws.append(hdr)
    for c in ws[1]: c.font = HD; c.fill = FILL; c.alignment = Alignment(horizontal="center", wrap_text=True)
    for r in rows: ws.append(r)
    for i, w in enumerate(widths or [], 1): ws.column_dimensions[ws.cell(1, i).column_letter].width = w
    ws.freeze_panes = "A2"; return ws
ws = wb.active; ws.title = "README"
for l in ["AILK014-007 / -008 / -009 — 750mer direct-print nanopore QC (Q-724749)", "Twist Bioscience R&D · 2026-09-27", "",
          f"{core['total_reads']:,} reads, three PromethION flowcells (one per sub-library), ONT barcoding off.",
          "Pairing pass: every read. Block-deletion / strand / full-length pass: every 8th read.", "",
          "DENOMINATORS:", "  *_pct_of_all       share of every read in fastq_pass",
          "  *_pct_of_assigned  share of reads with BOTH halves assigned (the pairing denominator)",
          "  deletion rates     share of reads carrying BOTH primer sites and assigned to a variant (1-in-8 subsample)",
          "  full-length length  primer-to-primer span after cutadapt (ONT adapters excluded); raw FASTQ length is still used for concatamer diagnostics", "",
          "BOUNDS: mosaic rate is a LOWER bound; intact fractions are UPPER bounds; screening-depth cell counts are MINIMA.",
          "Dropout is measured (rarefaction flat from 5% of depth), not a bound.", "",
          f"CONTROLS: 0 false chimeras in 300,000 simulated clean reads at 1–3% error (95% UB 0.006% per condition); 100% within-well swap detection.",
          "Cross-well reads are mostly ligation-fused (≥1,200 nt); they are kept in the conservative total but are not pairing chimera."]:
    ws.append([l])
ws.column_dimensions["A"].width = 110
S = []
def add(m, f, note=""):
    try:
        S.append([m] + [f(o) for o in ORD] + [f(None) if f(None) is not None else "", note])
    except (KeyError, TypeError):
        S.append([m] + ["" for _ in ORD] + ["", note + " (not computed for this run)"])
g = lambda d, k: (lambda o: d[o][k] if o else None)
add("Reads", lambda o: ps[o]["n_reads"] if o else core["total_reads"], "all reads in fastq_pass")
add("Both halves assigned, % of all", lambda o: ps[o]["assigned_pct_of_all"] if o else ov["assigned_pct_of_all"])
add("Correctly paired, % of assigned", lambda o: ps[o]["clean_designed_pct_of_assigned"] if o else ov["clean_designed_pct_of_assigned"])
add("Within-well chimera, % of assigned", lambda o: ps[o]["within_well_pct_of_assigned"] if o else ov["within_well_pct_of_assigned"], "spec ≤5%")
add("Cross-well, % of assigned", lambda o: ps[o]["cross_well_pct_of_assigned"] if o else ov["cross_well_pct_of_assigned"], "mostly ligation-fused reads")
add("Cross-sub-library, % of assigned", lambda o: ps[o]["cross_sublib_pct_of_assigned"] if o else ov["cross_sublib_pct_of_assigned"], "separate flowcells: control / carry-over")
add("Total mispairing (conservative), % of assigned", lambda o: round(tm[o], 4) if o else round(TMall, 4), "spec ≤5%")
add("Within-well, single-molecule reads only, %", lambda o: ext["single_molecule_pairing_per_sublib"][o]["within_well_pct_of_assigned"] if o else sm["within_well_pct_of_assigned"])
add("One half only, % of all", lambda o: ps[o]["one_half_pct_of_all"] if o else ov["one_half_pct_of_all"])
add("Unmapped, % of all", lambda o: ps[o]["unmapped_pct_of_all"] if o else ov["unmapped_pct_of_all"])
add("≥50 nt block deletion, % of both-primer reads", lambda o: DEL[o] if o else SC["blockdel50_pct_of_eligible"], "no spec today")
add("≥50 nt block deletion, correctly paired both-primer reads", lambda o: BD["per_sublib"][o]["pct"] if o else BD["block_deletion"]["50"]["pct"], "comparable to AILK013 per-order figures")
add("Intact (primer-span ≥90% AND no deletion), both-primer, %", lambda o: SC["fl_readlen_nogap_by_sublib_both"][o] if o else SC["fl_matrix"]["readlen_and_nogap"]["both_primer"], "spec ≥60%; length is primer-to-primer, not raw FASTQ")
add("Variants designed", lambda o: U[o]["n_designed"] if o else U["ALL"]["n_designed"])
add("Variant dropout", lambda o: U[o]["dropout"] if o else U["ALL"]["dropout"], "measured")
add("Reads/variant median", lambda o: U[o]["median"] if o else U["ALL"]["median"])
add("Reads/variant min", lambda o: U[o]["min"] if o else U["ALL"]["min"])
add("CV (zeros included)", lambda o: round(U[o]["cv"], 4) if o else round(U["ALL"]["cv"], 4), "spec ≤1.0")
add("Gini", lambda o: round(U[o]["gini"], 4) if o else round(U["ALL"]["gini"], 4))
add("95/5 ratio", lambda o: (round(U[o]["p95p5"], 3) if U[o].get("p95p5") else None) if o
    else (round(U["ALL"]["p95p5"], 3) if U["ALL"].get("p95p5") else None), "spec ≤20")
add("Within 2× of median, %", lambda o: round(U[o]["within2x"], 2) if o else round(U["ALL"]["within2x"], 2), "spec ≥60%")
add("Cells for 95% recovery (counting deletions, minimum)", lambda o: REC[o]["95"]["counting_deletions"] if o else None, "per sub-library pool")
if dsg:
    add("Internal fwd-primer 3' site at nt 175-195, % constructs", lambda o: dsg[o]["pct_site_175_195"] if o else None, "design screen")
    add("Deletion events at that site, %", lambda o: bd["primer_site_attribution"][o]["pct"] if o else None)
sheet("Summary", ["Metric"] + ORD + ["All", "Note"], S, [52, 14, 14, 14, 14, 44])
sheet("Controls", ["Caller", "Control", "Condition", "n", "hits", "%", "95% UB if zero"],
      [["classifier", "negative (false chimera)", k, v["n"], v["false_chimera"], v["pct"], v["upper95_pct_if_zero"]]
       for k, v in CT["classifier"]["negative"].items()]
      + [["classifier", "positive (within-well swap)", k, v["n"], v["detected"], v["pct"], None]
         for k, v in CT["classifier"]["positive"].items()]
      + [["deletion", "negative (false >=50 nt call)", k, v["n_after_filter"], v["false_calls"], v["pct"], v["upper95_pct_if_zero"]]
         for k, v in CT.get("deletion_caller", {}).get("negative", {}).items()]
      + [["deletion", "positive (programmed deletion, nt)", k, v["n_after_filter"], v["detected"], v["sensitivity_pct"], None]
         for k, v in CT.get("deletion_caller", {}).get("positive", {}).items()]
      + [["deletion", "both-primer filter", k, v, None, None, None]
         for k, v in CT.get("deletion_caller", {}).get("truncation_filter", {}).items()],
      [12, 26, 22, 10, 8, 10, 14])
sheet("Read_length_by_class", ["Class", "n", "median", "% 600–900 nt", "% ≥1,200 nt"],
      [[k, v["n"], v["median"], v["pct_single"], v["pct_fused"]] for k, v in ext["length_classes"].items()], [18, 14, 10, 14, 14])
_fln = {"aln_span_ge90": "alignment span ≥90% of design",
        "readlen_ge90": "primer-span length ≥90% of design",
        "no_gap50": "no internal deletion ≥50 nt",
        "readlen_and_nogap": "primer-span ≥90% AND no deletion ≥50 nt"}
pspan, psub = SC.get("primer_span", {}), SC.get("primer_span_by_sublib", {})
if pspan:
    ga = pspan.get("gene_assigned", {})
    prow = [["all gene-assigned both-primer", ga.get("n"), ga.get("median"), ga.get("p5"), ga.get("p95"),
             pspan.get("pct_within_30nt_of_design"), pspan.get("n_span_fail")]]
    for o in ORD:
        d = psub.get(o, {})
        prow.append([o, d.get("n"), d.get("median"), d.get("p5"), d.get("p95"), None, None])
    sheet("Primer_span_length",
          ["group", "n", "median nt", "p5", "p95", "% within 30 nt of design", "span trim failures"],
          prow, [28, 12, 12, 8, 8, 22, 18])
sheet("Block_deletion", ["Metric", "Value"],
      [["subsampled reads", SC["n_subsampled"]], ["both-primer %", SC["both_primer_pct"]], ["eligible (both primers, gene assigned)", SC["eligible_any_cls_n"]],
       ["≥50 nt calls", SC["blockdel50_n"]], ["rate %", SC["blockdel50_pct_of_eligible"]], ["median size nt", bd["size_median"]],
       ["IQR", f"{bd['size_iqr'][0]:.0f}–{bd['size_iqr'][1]:.0f}"], ["single-event reads %", round(bd["one_gap_pct"], 2)],
       ["edlib confirmation %", BD["edlib_confirm_pct"]]] +
      [[f"rate at ≥{k} nt (correctly paired) %", v["pct"]] for k, v in BD["block_deletion"].items()] +
      [[f"{o} per-well min/median/max %", fmt3(bd["per_well"].get(o, {}), "min", "median", "max")] for o in ORD] +
      [[f"{o} per-variant p5/median/p95 %", fmt3(bd["per_gene_spread"].get(o, {}), "p5", "median", "p95")] for o in ORD], [48, 30])
sheet("Full_length_definitions", ["Definition (% full length)", "Gene-assigned reads", "Both-primer reads", "Correctly paired reads"],
      [[_fln.get(k, k), v.get("gene_assigned"), v.get("both_primer"), v.get("clean_designed")] for k, v in SC["fl_matrix"].items()], [40, 16, 16, 16])
sheet("Screening_depth", ["Sub-library", "Target %", "Ignoring deletions", "Counting deletions", "Ratio", "Mass-weighted intact"],
      [[o, t, REC[o][t]["ignoring_deletions"], REC[o][t]["counting_deletions"], REC[o][t]["ratio"], round(REC[o]["mass_weighted_intact"], 4)] for o in ORD for t in ("90", "95", "99")], [14, 10, 18, 18, 8, 18])
pg = {r["gene"]: int(r["n"]) for r in csv.DictReader(open(D + "per_gene.csv"))}
pga = {(r["sublib"], r["gene"]): (int(r["n"]), int(r["g50"])) for r in csv.DictReader(open(D + "blockdel_per_gene_all.csv"))}
rows = []
for gname in sorted(meta):
    m = meta[gname]; n, d = pga.get((m["order"], gname), (0, 0))
    rows.append([gname, m["order"], m.get("framework", ""), m.get("pool", ""), m["well"], m.get("well_rep2", ""),
                 m["full_len"], pg.get(gname, 0), n, d, round(100 * d / n, 2) if n else None])
sheet("Per_variant", ["gene", "order", "framework", "pool", "well_rep1", "well_rep2", "length", "correctly_paired_reads", "subsample_both_primer_reads", "subsample_del50", "del50_pct"], rows, [26, 13, 10, 9, 9, 9, 8, 12, 12, 10, 9])
mh = {(r["order"], r["well"]): r for r in csv.DictReader(open(A.microhomology))} if A.microhomology else {}
rows = []
for r in csv.DictReader(open(D + "per_well.csv")):
    k = (r["order"], r["well"]); m = mh.get(k, {}); na = int(r["n_assigned"])
    rows.append([r["order"], r["well"], na, int(r["n_clean"]), int(r["n_within"]), int(r["n_cross"]), round(100 * int(r["n_within"]) / na, 4),
                 round(100 * int(r["n_cross"]) / na, 4), bd["per_well"][r["order"]]["_wells"].get(r["well"]), m.get("n_mates"), m.get("shared_18mers"), m.get("max_lcs")])
sheet("Per_well", ["order", "well", "n_assigned", "n_clean", "n_within", "n_cross", "within_pct", "cross_pct", "del50_pct", "n_mates", "shared_18mers", "max_lcs"], rows)
if core.get("rarefaction"):
    sheet("Rarefaction", ["fraction", "clean reads", "distinct variants"],
          [[r["frac"], r["n_reads"], r["n_variants"]] for r in core["rarefaction"]])
wb.save(f"{O}{A.name}_nanopore_QC.xlsx"); print("wrote", f"{O}{A.name}_nanopore_QC.xlsx")
r95 = {o: REC[o]["95"]["counting_deletions"] for o in ORD}
summary = [
    f"# {A.name} — nanopore QC summary",
    "",
    f"{core['total_reads']:,} reads in {core['n_files']} FASTQ files; {len(meta):,} designed constructs in "
    + ", ".join(f"{o} ({U[o]['n_designed']:,})" for o in ORD) + ".",
    "",
    "| Metric | " + " | ".join(ORD) + " |",
    "|---|" + "---|" * len(ORD),
    "| Reads | " + " | ".join(f"{ps[o]['n_reads']:,}" for o in ORD) + " |",
    "| Correctly paired (% of assigned) | " + " | ".join(f"{ps[o]['clean_designed_pct_of_assigned']:.2f}%" for o in ORD) + " |",
    "| Within-well chimera (% of assigned) | " + " | ".join(f"{ps[o]['within_well_pct_of_assigned']:.4f}%" for o in ORD) + " |",
    "| Total mispairing, conservative | " + " | ".join(f"{tm[o]:.3f}%" for o in ORD) + " |",
    "| >=50 nt block deletion (both-primer reads) | " + " | ".join(f"{DEL[o]:.2f}%" if DEL[o] is not None else "n/a" for o in ORD) + " |",
    "| Intact full-length (primer-span, both-primer) | " + " | ".join(
        (lambda v: f"{v:.1f}%" if v is not None else "n/a")(SC["fl_readlen_nogap_by_sublib_both"].get(o)) for o in ORD) + " |",
    "| Variant dropout | " + " | ".join(f"{U[o]['dropout']} / {U[o]['n_designed']:,}" for o in ORD) + " |",
    "| Reads per variant, CV | " + " | ".join(f"{U[o]['cv']:.3f}" for o in ORD) + " |",
    "| Reads per variant, 95/5 | " + " | ".join(f"{U[o]['p95p5']:.2f}" if U[o].get("p95p5") else "n/a" for o in ORD) + " |",
    "| Cells for 95% recovery (minimum) | " + " | ".join(f"{r95[o]:,}" if r95[o] else "n/a" for o in ORD) + " |",
    "",
    "## How to read these numbers",
    "",
    "- Pairing percentages are shares of reads with **both halves assigned**, not of all reads.",
    f"- Deletion rates are shares of reads carrying **both primer sites** and assigned to a variant, on a 1-in-{BD['subsample']} subsample.",
    "- Intact fractions use **primer-to-primer span** (not raw FASTQ length) among reads carrying both primer sites, and are **upper bounds** (the caller cannot see deletions below 50 nt or substitutions); screening-depth cell counts are therefore **minima**.",
    "- Cross-well reads are largely ligation-fused molecules from sequencing prep, not chimeric constructs — see the read-length-by-class sheet before quoting the conservative total.",
    "",
    "## Calibration",
    "",
]
neg = CT["classifier"]["negative"]
zero = [k for k, v in neg.items() if v["false_chimera"] == 0]
summary += [
    f"- Classifier: {sum(neg[k]['n'] for k in zero):,} simulated clean reads across {len(zero)} error conditions with zero false chimera calls "
    f"(95% upper bound {max(neg[k]['upper95_pct_if_zero'] for k in zero):.3f}% per condition); "
    + ", ".join(f"{k} {v['pct']:.3f}%" for k, v in neg.items() if v["false_chimera"]) or "no condition produced a false call",
    "- Classifier sensitivity to injected within-well swaps: "
    + ", ".join(f"{k} {v['pct']:.1f}%" for k, v in CT["classifier"]["positive"].items()) + ".",
]
dc = CT.get("deletion_caller", {})
if dc:
    summary.append("- Deletion caller: "
                   + ", ".join(f"{k} nt {v['sensitivity_pct']:.1f}%" for k, v in dc.get("positive", {}).items())
                   + " sensitivity; false calls "
                   + ", ".join(f"{k} {v['pct']:.2f}%" for k, v in dc.get("negative", {}).items()) + ".")
open(f"{O}{A.name}_qc_summary.md", "w").write("\n".join(summary) + "\n")
print("wrote", f"{O}{A.name}_qc_summary.md")
